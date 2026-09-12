"""
Generate realistic sample Excel datasets for automated reconciliation demonstration.
Creates:
- data/raw_input.xlsx
- data/sources/source_1_erp.xlsx
- data/sources/source_2_procurement.xlsx
- data/sources/source_3_warehouse.xlsx
- data/sources/source_4_finance.xlsx
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_styled_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin = Side(border_style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[1].height = 24

    # Rows
    for r_idx, row_vals in enumerate(rows, start=2):
        ws.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(val, float):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column widths
    for col in ws.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    return ws

def generate_samples(base_dir="data"):
    os.makedirs(os.path.join(base_dir, "sources"), exist_ok=True)

    # 1. RAW INPUT WORKBOOK
    raw_wb = openpyxl.Workbook()
    raw_wb.remove(raw_wb.active)  # remove default sheet

    invoices_headers = ["Invoice_ID", "Vendor_Name", "Invoice_Amount", "Tax_Amount", "Currency", "Payment_Status"]
    raw_invoices = [
        ["INV-2024-001", "Global Logistics Ltd", 15400.00, 1540.00, "USD", "Approved"],
        ["INV-2024-002", "Apex Cloud Solutions", 8250.00, 825.00, "USD", "Paid"],
        ["INV-2024-003", "Nexus Office Supply", 3410.50, 341.05, "USD", "Pending"],
        ["INV-2024-004", "Quantum Dynamics Corp", 42000.00, 4200.00, "USD", "Approved"],
        ["INV-2024-005", "Starlight Media Group", 6750.00, 675.00, "USD", "Pending"],
        ["INV-2024-006", "Vanguard Logistics Inc", 19800.00, 1980.00, "USD", "Approved"],
        ["INV-2024-007", "Beacon Energy Services", 12300.00, 1230.00, "USD", "Paid"],
        ["INV-2024-008", "Horizon Cyber Systems", 27500.00, 2750.00, "USD", "Approved"],
        ["INV-2024-009", "Pinnacle Hardware Inc", 5840.00, 584.00, "USD", "Pending"],
        ["INV-2024-010", "Aura Creative Agency", 9100.00, 910.00, "USD", "Paid"],
    ]
    create_styled_sheet(raw_wb, "Invoices", invoices_headers, raw_invoices)

    inventory_headers = ["SKU", "Item_Description", "Unit_Cost", "Stock_Qty", "Warehouse_Zone"]
    raw_inventory = [
        ["SKU-101", "Industrial Server Rack 42U", 1250.00, 45, "Zone-A"],
        ["SKU-102", "Cat6A Shielded Ethernet Cable 305m", 185.50, 120, "Zone-B"],
        ["SKU-103", "Gigabit Enterprise Switch 48-Port", 890.00, 30, "Zone-A"],
        ["SKU-104", "Uninterruptible Power Supply 3000VA", 1450.00, 18, "Zone-C"],
        ["SKU-105", "Fiber Optic Transceiver 10G SFP+", 95.00, 250, "Zone-B"],
    ]
    create_styled_sheet(raw_wb, "Inventory", inventory_headers, raw_inventory)

    raw_path = os.path.join(base_dir, "raw_input.xlsx")
    raw_wb.save(raw_path)

    # 2. SOURCE 1: ERP Ledger (Has exact matches, 1 typo, 1 small 2% variance)
    src1_wb = openpyxl.Workbook()
    src1_wb.remove(src1_wb.active)
    src1_invoices = [
        ["INV-2024-001", "Global Logistics Ltd", 15400.00, 1540.00, "USD", "Approved"],  # Exact
        ["INV-2024-002", "Apex Cloud Solutons", 8250.00, 825.00, "USD", "Paid"],       # Close text (typo Solutons >80%) -> YELLOW
        ["INV-2024-003", "Nexus Office Supply", 3480.00, 348.00, "USD", "Pending"],     # 2% numeric diff -> YELLOW
        ["INV-2024-004", "Quantum Dynamics Corp", 42000.00, 4200.00, "USD", "Approved"], # Exact
        ["INV-2024-005", "Starlight Media Group", 6750.00, 675.00, "USD", "Pending"],   # Exact
    ]
    create_styled_sheet(src1_wb, "Invoices", invoices_headers, src1_invoices)
    src1_path = os.path.join(base_dir, "sources", "source_1_erp.xlsx")
    src1_wb.save(src1_path)

    # 3. SOURCE 2: Procurement Batch (Has exact, 1 mismatch in amount, 1 missing status)
    src2_wb = openpyxl.Workbook()
    src2_wb.remove(src2_wb.active)
    src2_invoices = [
        ["INV-2024-006", "Vanguard Logistics Inc", 19800.00, 1980.00, "USD", "Approved"], # Exact
        ["INV-2024-007", "Beacon Energy Services", 19500.00, 1950.00, "USD", "Paid"],     # Mismatch (58% diff) -> RED
        ["INV-2024-008", "Horizon Cyber Systems", 27500.00, 2750.00, "USD", None],        # Empty/Null -> RED
        ["INV-2024-009", "Pinnacle Hardwre Inc", 5840.00, 584.00, "USD", "Pending"],     # Close text (>80%) -> YELLOW
        ["INV-2024-010", "Aura Creative Agency", 9100.00, 910.00, "USD", "Paid"],        # Exact
    ]
    create_styled_sheet(src2_wb, "Invoices", invoices_headers, src2_invoices)
    src2_path = os.path.join(base_dir, "sources", "source_2_procurement.xlsx")
    src2_wb.save(src2_path)

    # 4. SOURCE 3: Warehouse Physical Count
    src3_wb = openpyxl.Workbook()
    src3_wb.remove(src3_wb.active)
    src3_inventory = [
        ["SKU-101", "Industrial Server Rack 42U", 1250.00, 45, "Zone-A"],              # Exact
        ["SKU-102", "Cat6A Shielded Ethernet Cable 305m", 185.50, 122, "Zone-B"],       # 122 vs 120 (1.6% diff) -> YELLOW
        ["SKU-103", "Gigabit Enterprise Switch 48-Port", 890.00, 30, "Zone-A"],        # Exact
        ["SKU-104", "Uninterruptible Power Supply 3000VA", 1450.00, 8, "Zone-C"],       # 8 vs 18 (55% diff) -> RED
        ["SKU-105", "Fiber Optic Transceiver 10G SFP+", 95.00, None, "Zone-B"],        # Null stock -> RED
    ]
    create_styled_sheet(src3_wb, "Inventory", inventory_headers, src3_inventory)
    src3_path = os.path.join(base_dir, "sources", "source_3_warehouse.xlsx")
    src3_wb.save(src3_path)

    # 5. SOURCE 4: Finance Audit Cross-Check
    src4_wb = openpyxl.Workbook()
    src4_wb.remove(src4_wb.active)
    src4_invoices = [
        ["INV-2024-001", "Global Logistics Ltd", 15400.00, 1540.00, "USD", "Approved"],
        ["INV-2024-004", "Quantum Dynamics Corp", 42000.00, 4200.00, "USD", "Approved"],
    ]
    create_styled_sheet(src4_wb, "Invoices", invoices_headers, src4_invoices)
    src4_path = os.path.join(base_dir, "sources", "source_4_finance.xlsx")
    src4_wb.save(src4_path)

    print(f"Sample Excel workbooks successfully generated in '{base_dir}' directory!")
    return {
        "raw": raw_path,
        "sources_dir": os.path.join(base_dir, "sources"),
        "source_files": [src1_path, src2_path, src3_path, src4_path]
    }

if __name__ == "__main__":
    generate_samples()
