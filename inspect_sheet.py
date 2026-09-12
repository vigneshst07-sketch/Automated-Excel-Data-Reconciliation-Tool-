"""
Quick helper to inspect an Excel workbook and return cells with OpenPyXL fill colors in JSON format.
"""

import sys
import json
import argparse
import openpyxl

def inspect_sheet(file_path: str, sheet_name: str = None, max_rows: int = 100, max_cols: int = 20):
    wb = openpyxl.load_workbook(file_path, data_only=False)
    sheet_names = wb.sheetnames
    target_sheet = sheet_name if sheet_name and sheet_name in sheet_names else sheet_names[0]
    ws = wb[target_sheet]

    max_r = min(ws.max_row, max_rows)
    max_c = min(ws.max_column, max_cols)

    rows = []
    for r in range(1, max_r + 1):
        row_data = []
        for c in range(1, max_c + 1):
            cell = ws.cell(row=r, column=c)
            val = str(cell.value) if cell.value is not None else ""

            # Detect OpenPyXL fill color
            fill_type = "NORMAL"
            if cell.fill and cell.fill.start_color:
                rgb = str(cell.fill.start_color.rgb or "")
                if "FFEB9C" in rgb or "EB9C" in rgb:
                    fill_type = "YELLOW"
                elif "FFC7CE" in rgb or "C7CE" in rgb:
                    fill_type = "RED"
                elif "0F172A" in rgb or "1E293B" in rgb:
                    fill_type = "HEADER_DARK"
                elif "F1F5F9" in rgb or "F8FAFC" in rgb:
                    fill_type = "HEADER_LIGHT"

            row_data.append({
                "val": val,
                "fill": fill_type,
                "coord": cell.coordinate,
                "row": r,
                "col": c
            })
        rows.append(row_data)

    return {
        "sheet_names": sheet_names,
        "active_sheet": target_sheet,
        "rows": rows
    }

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--sheet", default=None)
    args = parser.parse_args()

    result = inspect_sheet(args.file, args.sheet)
    print(json.dumps(result))
